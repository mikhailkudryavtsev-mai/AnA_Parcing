import requests
from bs4 import BeautifulSoup
import csv
import re
import math
import time
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches

BASE_URL = "https://mai.ru/press/news/"
TAGS = "10312"  # Образование 2.0
CATEGORY_NAME = "Образование 2.0"

CSV_FILE = "mai_news.csv"
PNG_FILE = "mai_news_chart.png"
DOCX_FILE = "mai_news_report.docx"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                  "AppleWebKit/537.36 Chrome/120.0 Safari/537.36"
}

MONTHS = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4,
    "май": 5, "июн": 6, "июл": 7, "авг": 8,
    "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
}

MONTH_NAMES = {
    1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
    5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
    9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь",
}


def get_page_url(page):
    if page == 1:
        return f"{BASE_URL}?tags={TAGS}"
    return f"{BASE_URL}?tags={TAGS}&PAGEN_1={page}"


def get_total_pages(soup):
    pagination_block = soup.select_one("nav.pagination, .pagination")
    if pagination_block:
        parent = pagination_block.find_parent()
        if parent:
            text = parent.get_text()
            match = re.search(r"из\s+(\d+)", text)
            if match:
                total = int(match.group(1))
                return math.ceil(total / 16) 

    max_page = 1
    for a in soup.select("a.page-link"):
        href = a.get("href", "")
        match = re.search(r"PAGEN_1=(\d+)", href)
        if match:
            max_page = max(max_page, int(match.group(1)))
    return max_page


def parse_page(soup):
    items = []
    cards = soup.select("div.col-sm-6.col-lg-6.mb-3.mb-lg-5")

    for card in cards:

        h5 = card.select_one("div.card-body h5")
        title = h5.get_text(strip=True) if h5 else ""

        badge = card.select_one("span.badge")
        date_text = badge.get_text(strip=True) if badge else ""

        link_tag = card.select_one("a.card")
        link = ""
        if link_tag:
            href = link_tag.get("href", "")
            link = f"https://mai.ru/press/news/{href}" if href else ""

        if title:
            items.append({
                "title": title,
                "date_text": date_text,
                "link": link,
            })

    return items


def parse_date(date_text):
    parts = date_text.split()
    if len(parts) >= 2:
        try:
            day = int(parts[0])
            month = MONTHS.get(parts[1].lower(), 0)
            if month:
                return day, month
        except ValueError:
            pass
    return None, None


def assign_years(items):
    current_year = datetime.now().year
    prev_month = None

    for item in items:
        day, month = parse_date(item["date_text"])
        if month:
            if prev_month is not None and month > prev_month:
                current_year -= 1
            prev_month = month
            item["date"] = f"{day:02d}.{month:02d}.{current_year}"
            item["year"] = current_year
            item["month"] = month
        else:
            item["date"] = ""
            item["year"] = None
            item["month"] = None


def save_csv(items):
    with open(CSV_FILE, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["title", "date", "year", "month", "link"])
        writer.writeheader()
        for item in items:
            writer.writerow({
                "title": item["title"],
                "date": item["date"],
                "year": item.get("year", ""),
                "month": item.get("month", ""),
                "link": item["link"],
            })
    print(f"CSV сохранён: {CSV_FILE} ({len(items)} записей)")

def build_chart(items):
    df = pd.DataFrame(items)
    df = df[df["year"].notna()].copy()
    df["year"] = df["year"].astype(int)
    df["month"] = df["month"].astype(int)

    grouped = df.groupby(["year", "month"]).size().reset_index(name="count")
    grouped = grouped.sort_values(["year", "month"])

    grouped["label"] = grouped.apply(
        lambda r: f"{MONTH_NAMES[int(r['month'])][:3]} {int(r['year'])}", axis=1
    )

    plt.rcParams["font.family"] = "DejaVu Sans"
    fig, ax = plt.subplots(figsize=(max(8, len(grouped) * 0.6), 5))

    bars = ax.bar(grouped["label"], grouped["count"], color="steelblue", edgecolor="white")

    for bar in bars:
        height = bar.get_height()
        ax.text(bar.get_x() + bar.get_width() / 2., height,
                int(height), ha="center", va="bottom", fontsize=9)

    ax.set_title(f"Распределение новостей по месяцам\n(рубрика «{CATEGORY_NAME}»)", fontsize=14)
    ax.set_xlabel("Месяц")
    ax.set_ylabel("Количество новостей")
    plt.xticks(rotation=45, ha="right")
    plt.tight_layout()
    plt.savefig(PNG_FILE, dpi=150)
    plt.close()

    print(f"График сохранён: {PNG_FILE}")
    return grouped


def build_docx(items, monthly):
    doc = Document()
    doc.add_heading("Отчёт по новостям МАИ", 0)

    doc.add_paragraph(f"Рубрика: {CATEGORY_NAME}")
    doc.add_paragraph(f"Источник: {BASE_URL}?tags={TAGS}")
    doc.add_paragraph(f"Всего публикаций: {len(items)}")
    doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

    doc.add_heading("Распределение по месяцам", level=1)
    doc.add_picture(PNG_FILE, width=Inches(6))

    doc.add_heading("Публикации", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    hdr[0].text = "Дата"
    hdr[1].text = "Заголовок"
    hdr[2].text = "Ссылка"

    for item in items:
        row = table.add_row().cells
        row[0].text = item["date"]
        row[1].text = item["title"]
        row[2].text = item["link"]

    doc.save(DOCX_FILE)
    print(f"✓ DOCX отчёт сохранён: {DOCX_FILE}")


def main():
    session = requests.Session()
    session.headers.update(HEADERS)

    print(f"Загружаем: {get_page_url(1)}")
    resp = session.get(get_page_url(1), timeout=15)
    resp.encoding = "utf-8"
    soup = BeautifulSoup(resp.text, "html.parser")

    total_pages = get_total_pages(soup)
    print(f"Всего страниц: {total_pages}")

    all_items = parse_page(soup)
    print(f"  Стр. 1: {len(all_items)} новостей")

    for page in range(2, total_pages + 1):
        time.sleep(0.5)
        url = get_page_url(page)
        resp = session.get(url, timeout=15)
        resp.encoding = "utf-8"
        soup = BeautifulSoup(resp.text, "html.parser")
        items = parse_page(soup)
        all_items.extend(items)
        print(f"  Стр. {page}: {len(items)} новостей")

    print(f"\nИтого: {len(all_items)} новостей")

    assign_years(all_items)

    save_csv(all_items)

    monthly = build_chart(all_items)

    build_docx(all_items, monthly)


if __name__ == "__main__":
    main()
